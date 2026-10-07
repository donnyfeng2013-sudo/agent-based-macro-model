"""
Consumption_analysis.py: calibrate household consumption from the BLS Consumer Expenditure Survey 2024.

Run once from anywhere:  python3 Model/Consumption_analysis.py
Reads   Data/BLS CES 2024.xlsx   (Table 1203: income before taxes, 9 income bands)
Writes  Data/Calibration/
    consumption_by_band.csv     per income band: consumer units, mean income before taxes,
                                mean consumption (USD) and its standard error
    consumption_function.csv    C = C_0 + a × Y fitted across bands: C_0, a and their standard errors
    income_elasticity.csv       income elasticity of each model sector (ABM2.jl SECTORS), from
                                CES categories mapped to sectors; unmapped sectors get 1

Consumption = average annual expenditures minus items that are not purchases from firms
(EXCLUDE below). Each band reports a mean income and a mean consumption, so the regression
uses one point per band, weighted by the number of consumer units in it.
"""

import re
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

# ------------------------------------------------------------------
# Settings
# ------------------------------------------------------------------
ROOT     = Path(__file__).resolve().parent.parent
CES_FILE = ROOT / "Data" / "BLS CES 2024.xlsx"
OUT_DIR  = ROOT / "Data" / "Calibration"

EXCLUDE  = ["Personal insurance and pensions",   # mostly Social Security contributions: a tax, not a purchase
            "Cash contributions"]                # gifts and donations to other households / charities
MIN_INCOME = 0          # drop bands with mean income below this (USD); 0 keeps all 9 bands

# ------------------------------------------------------------------
# 1. Load Table 1203
# ------------------------------------------------------------------
STATS = {"Mean", "Share", "SE", "RSE"}

def load_ces(path=CES_FILE):
    """Return {item path: {stat: array over bands}} and the band labels.
    Item path joins the labels of an item and its parents (from Excel indentation),
    e.g. 'Food > Food at home', so repeated labels stay distinct."""
    ws = openpyxl.load_workbook(path, data_only=True).active
    header = next(r for r in range(1, 20) if ws.cell(r, 1).value == "Item")
    n_cols = ws.max_column
    bands  = [" ".join(str(ws.cell(header, c).value).split()) for c in range(2, n_cols + 1)]

    table, stack, current = {}, [], None              # stack holds (indent, label) of open parents
    for r in range(header + 1, ws.max_row + 1):
        label = ws.cell(r, 1).value
        if label is None:
            continue
        label = re.sub(r"(\s+[a-z]/)+\s*$", "", str(label)).strip()   # drop footnote markers, e.g. ' a/'
        values = [ws.cell(r, c).value for c in range(2, n_cols + 1)]
        if label in STATS:
            if current is not None:
                table[current][label] = np.array([v if isinstance(v, (int, float)) else np.nan for v in values], float)
            continue
        indent = ws.cell(r, 1).alignment.indent or 0
        while stack and stack[-1][0] >= indent:
            stack.pop()
        stack.append((indent, label))
        current = " > ".join(l for _, l in stack)
        table.setdefault(current, {})
        if isinstance(values[0], (int, float)):        # one-line items, e.g. number of consumer units
            table[current]["Mean"] = np.array(values, float)
    return table, bands

def item(table, label):
    """Statistics of the item whose own label is `label` (first match, top of table first)."""
    for key, stats in table.items():
        if key.split(" > ")[-1] == label and stats:
            return stats
    raise KeyError(label)

# ------------------------------------------------------------------
# 2. Consumption by income band
# ------------------------------------------------------------------
def consumption_by_band(table, bands):
    """Drop the 'All consumer units' column; one row per income band."""
    units  = item(table, "Number of consumer units (in thousands)")["Mean"]
    income = item(table, "Income before taxes")
    total  = item(table, "Average annual expenditures")
    excl   = sum(item(table, e)["Mean"] for e in EXCLUDE)
    df = pd.DataFrame({
        "band":              bands,
        "consumer_units":    units * 1000,
        "income":            income["Mean"],
        "income_se":         income["SE"],
        "consumption":       total["Mean"] - excl,
        "consumption_se":    total["SE"],    # SE of total expenditures; BLS gives no SE for the difference
        "earners":           item(table, "Earners")["Mean"],
    })
    df["apc"] = df.consumption / df.income   # average propensity to consume
    return df.iloc[1:].reset_index(drop=True)

# ------------------------------------------------------------------
# 3. Consumption function: C = C_0 + a × Y
# ------------------------------------------------------------------
def weighted_line(x, y, w):
    """Weighted least squares of y on [1, x]; returns (coef, standard errors).
    Standard errors are heteroskedasticity-robust (HC1) on the band means."""
    X  = np.column_stack([np.ones_like(x), x])
    W  = w / w.mean()
    XtWX_inv = np.linalg.inv(X.T @ (X * W[:, None]))
    coef  = XtWX_inv @ X.T @ (W * y)
    resid = y - X @ coef
    n, k  = X.shape
    meat  = X.T @ (X * (W**2 * resid**2)[:, None])
    cov   = XtWX_inv @ meat @ XtWX_inv * n / (n - k)
    return coef, np.sqrt(np.diag(cov))

def consumption_function(bands):
    b = bands[bands.income >= MIN_INCOME]
    coef, se = weighted_line(b.income.values, b.consumption.values, b.consumer_units.values)
    lowest = b.iloc[0]
    return pd.DataFrame({
        "parameter": ["C_0", "a", "C_0_lowest_band"],
        "value":     [coef[0], coef[1], lowest.consumption],
        "se":        [se[0], se[1], lowest.consumption_se],
        "note":      ["regression intercept (USD per consumer unit per year)",
                      "slope: extra consumption per extra USD of income before taxes",
                      "mean consumption of the lowest income band (alternative C_0)"],
    })

# ------------------------------------------------------------------
# 4. Income elasticities by CES category, mapped to model sectors
# ------------------------------------------------------------------
# Elasticity of category spending with respect to total consumption (the model's budget E):
# slope of log(category spending) on log(consumption) across the 9 bands, weighted by consumer units.

# model sectors (same order as SECTORS in ABM2.jl)
SECTORS = [
    "111", "112", "115", "211", "212", "213", "221", "236", "237", "238",
    "311", "312", "313", "314", "315", "316", "321", "322", "323", "324",
    "325", "326", "327", "331", "332", "333", "334", "335", "3361MV", "3364OT",
    "337", "339", "423", "424", "441", "444", "445", "449", "456", "457",
    "458", "459", "481", "484", "486", "488", "492", "512", "513", "517",
    "518", "519", "522", "523", "531", "532", "533", "5415", "5412OP", "561",
    "562", "611", "621", "713", "721", "722", "811", "812",
]

# sector => CES items households buy from it. A sector with several items takes their
# spending-weighted average elasticity. Sectors not listed get elasticity 1.
SECTOR_CES = {
    "111":    ["Fresh fruits", "Fresh vegetables"],                                     # crop production
    "112":    ["Eggs"],                                                                 # animal production
    "221":    ["Natural gas", "Electricity", "Water and other public services"],        # utilities
    "236":    ["Owned dwellings"],                                                      # construction of buildings
    "238":    ["Maintenance, repairs, insurance, and other expenses"],                  # specialty trade contractors (home repairs)
    "311":    ["Cereals and bakery products", "Meats, poultry, fish, and eggs", "Dairy products",
               "Processed fruits and vegetables", "Sugar and other sweets", "Fats and oils",
               "Miscellaneous foods"],                                                  # food manufacturing
    "312":    ["Nonalcoholic beverages", "Alcoholic beverages",
               "Tobacco products and smoking supplies"],                                # beverages and tobacco
    "314":    ["Household textiles", "Floor coverings"],                                # textile products
    "315":    ["Men and boys", "Women and girls", "Children under 2"],                  # apparel
    "316":    ["Footwear"],                                                             # leather
    "322":    ["Other household products"],                                             # paper
    "324":    ["Gasoline and other fuels", "Fuel oil and other fuels"],                 # petroleum products
    "325":    ["Drugs", "Laundry and cleaning products"],                               # chemicals
    "334":    ["Audio and visual equipment and services"],                              # electronics
    "335":    ["Major appliances", "Small appliances and miscellaneous housewares"],    # appliances
    "3361MV": ["Cars and trucks, new"],                                                 # motor vehicles
    "337":    ["Furniture"],                                                            # furniture
    "339":    ["Toys, hobbies, and playground equipment", "Medical supplies"],          # misc. manufacturing
    "441":    ["Vehicle purchases (net outlay)"],                                       # vehicle dealers
    "445":    ["Food at home"],                                                         # food retailers
    "449":    ["Household furnishings and equipment"],                                  # furniture & electronics retailers
    "456":    ["Drugs", "Personal care products and services"],                         # health & personal care retailers
    "457":    ["Gasoline and other fuels"],                                             # gas stations
    "458":    ["Apparel and services"],                                                 # clothing retailers
    "459":    ["Pets", "Toys, hobbies, and playground equipment", "Reading",
               "Other entertainment supplies, equipment, and services"],                # sporting, hobby, book retailers
    "481":    ["Public and other transportation"],                                      # air transport
    "492":    ["Postage and stationery"],                                               # couriers
    "512":    ["Audio and visual equipment and services"],                              # film, music, streaming
    "513":    ["Reading"],                                                              # publishing
    "517":    ["Telephone services"],                                                   # telecoms
    "522":    ["Mortgage interest and charges", "Vehicle finance charges"],             # credit
    "531":    ["Rented dwellings"],                                                     # real estate
    "532":    ["Vehicle rental, leases, licenses, and other charges"],                  # rental and leasing
    "5412OP": ["Miscellaneous"],                                                        # legal, accounting, funeral etc.
    "561":    ["Other household expenses"],                                             # household services (cleaning, lawn, security)
    "611":    ["Education"],                                                            # education
    "621":    ["Medical services"],                                                     # ambulatory health care
    "713":    ["Fees and admissions"],                                                  # recreation
    "721":    ["Other lodging"],                                                        # accommodation
    "722":    ["Food away from home"],                                                  # restaurants
    "811":    ["Maintenance and repairs"],                                              # vehicle repair
    "812":    ["Personal care products and services"],                                  # personal services
}

def category_elasticity(table, bands, label):
    """(elasticity, all-units mean spending) of one CES item."""
    spend = item(table, label)["Mean"]
    coef, _ = weighted_line(np.log(bands.consumption.values), np.log(spend[1:]),
                            bands.consumer_units.values)
    return coef[1], spend[0]

def sector_elasticities(table, bands):
    rows = []
    for s in SECTORS:
        labels = SECTOR_CES.get(s, [])
        if not labels:
            rows.append((s, 1.0, False, ""))
            continue
        est = [category_elasticity(table, bands, l) for l in labels]
        eta = sum(e * w for e, w in est) / sum(w for _, w in est)   # spending-weighted average
        rows.append((s, eta, True, "; ".join(labels)))
    return pd.DataFrame(rows, columns=["sector", "income_elasticity", "mapped", "ces_items"])

# ------------------------------------------------------------------
# Main
# ------------------------------------------------------------------
def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    table, labels = load_ces()
    bands = consumption_by_band(table, labels)
    func  = consumption_function(bands)
    bands.to_csv(OUT_DIR / "consumption_by_band.csv", index=False, float_format="%.4f")
    func.to_csv(OUT_DIR / "consumption_function.csv", index=False, float_format="%.6f")
    eta = sector_elasticities(table, bands)
    eta.to_csv(OUT_DIR / "income_elasticity.csv", index=False, float_format="%.4f")

    pd.set_option("display.width", 140)
    print(f"Consumption = average annual expenditures - {', '.join(EXCLUDE)}\n")
    print(bands[["band", "consumer_units", "income", "consumption", "consumption_se", "apc", "earners"]]
          .round({"income": 0, "consumption": 0, "consumption_se": 0, "apc": 2}).to_string(index=False))
    print(f"\nC = C_0 + a × Y  (bands with income >= {MIN_INCOME:,}, weighted by consumer units)")
    print(func.round(4).to_string(index=False))
    fit = func.set_index("parameter").value
    resid = bands.consumption - (fit.C_0 + fit.a * bands.income)
    print("\nFitted - actual by band (USD):", (-resid).round(0).astype(int).tolist())
    m = eta[eta.mapped]
    print(f"\nIncome elasticities: {len(m)} sectors mapped to CES, {len(eta) - len(m)} unmapped (set to 1)")
    print(m[["sector", "income_elasticity", "ces_items"]].round(2).to_string(index=False))
    print(f"Negative elasticities: {(m.income_elasticity < 0).sum()}")

if __name__ == "__main__":
    main()
